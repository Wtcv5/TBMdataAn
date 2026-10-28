from pathlib import Path
import json
import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'TBMData/TBMData'
items = []
for d in sorted(DATA.iterdir()):
    if not d.is_dir():
        continue
    files = sorted(p for p in d.rglob('tbm_data*.csv') if 'common_using' not in p.parts and 'codes' not in p.parts)
    if not files:
        continue
    df = pd.read_csv(files[0], nrows=1000, encoding='utf-8-sig')
    cols = [c for c in df if any(s in c for s in ['转速', '刀盘速度', '刀盘转矩', '刀盘扭矩', '推进速度', '推进力', '总推力', '贯入度'])]
    item = dict(directory=d.name, files=len(files), bytes=sum(p.stat().st_size for p in files), columns=len(df.columns), selected={c:pd.to_numeric(df[c],errors='coerce').describe().to_dict() for c in cols})
    items.append(item)
    print(d.name, len(files), len(df.columns), cols, flush=True)
(ROOT/'analysis/source_preview.json').write_text(json.dumps(items,ensure_ascii=False,indent=2),encoding='utf-8')
for p in DATA.glob('*.xlsx'):
    w = openpyxl.load_workbook(p,read_only=True,data_only=True)
    for s in w:
        print('WORKBOOK',p.name,s.title,s.max_row,s.max_column,flush=True)
        for row in s.iter_rows(min_row=1,max_row=8,values_only=True):
            print(row[:20],flush=True)
    w.close()
